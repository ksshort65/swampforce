import csv,json,glob,collections,sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
B='/workspace/checkpoint-review/'
R=list(csv.DictReader(open(B+'verified-items.csv',encoding='utf-8')))
fields=list(R[0].keys())
order=['verdicts_scorecard.json','verdicts_ledger_resource.json','verdicts_respot.json','verdicts_essays.json','verdicts_media.json','verdicts_late.json','verdicts_sweep.json','verdicts_followup.json']
M={}
for f in order:
    try:
        for k_,v_ in json.load(open(B+'fv/'+f)).items(): M.setdefault(k_,{}).update(v_)
    except FileNotFoundError: pass
n=0
for r in R:
    v=M.get(r['Item_No'])
    if v:
        for k in ('Verdict','Best_Source_URL','Source_Type','Fix_Needed','Notes'):
            if k in v: r[k]=v[k]
        if 'Evidence_Level' in v: r['Evidence_Level']=v['Evidence_Level']
        n+=1
with open(B+'verified-items.csv','w',newline='',encoding='utf-8') as fh:
    w=csv.DictWriter(fh,fieldnames=fields); w.writeheader(); w.writerows(R)
wb=Workbook(); ws=wb.active; ws.title='verified-items'; ws.append(fields)
hf=PatternFill('solid',fgColor='FF1B2A41'); hfont=Font(color='FFFFFF',bold=True)
for c in ws[1]: c.fill=hf; c.font=hfont
fills={'Verified':'FFC6EFCE','Verified with correction needed':'FFFFEB9C','Unverified':'FFDDEBF7','False':'FFFFC7CE',
       'False, cut it':'FFFFC7CE','Cannot verify, cut it':'FFF4B084','Opinion':'FFE2D5F1'}
vi=fields.index('Verdict')+1
for r in R:
    ws.append([r.get(c,'') for c in fields])
    col=fills.get(r['Verdict'])
    if col: ws.cell(ws.max_row,vi).fill=PatternFill('solid',fgColor=col)
ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
for col in ws.columns:
    ml=min(55,max(len(str(c.value or '')) for c in col)); ws.column_dimensions[col[0].column_letter].width=max(12,ml+2)
wb.save(B+'verified-items.xlsx')
print('applied',n); print(collections.Counter(r['Verdict'] for r in R))
print(collections.Counter((r['Category'],r['Verdict']) for r in R if r['Verdict']=='Unverified'))
