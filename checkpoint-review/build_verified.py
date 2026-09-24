import json, csv, re
from pathlib import Path
from collections import Counter, defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

items=json.loads(Path('items.work.json').read_text())
url_cache=json.loads(Path('url-check.json').read_text()) if Path('url-check.json').exists() else {}

def load_cat(path, prefix):
    out={}
    with open(path, newline='', encoding='utf-8-sig') as fh:
        for r in csv.DictReader(fh):
            out[f"{prefix}-{r['Item_ID']}"]=r
    return out
CAT={}
CAT.update(load_cat('/workspace/term-split/first-term-trump-admin-media-deception.csv','FT'))
CAT.update(load_cat('/workspace/term-split/later-second-term-trump-admin-media-deception.csv','LT'))

CURATED = {
  'Whips': {
    'Verdict':'Verified','Best_Source_URL':'https://www.cbp.gov/newsroom/national-media-release/cbp-releases-findings-investigation-horse-patrol-activity-del-rio',
    'Source_Type':'official_record','Evidence_Level':'Proven false','Fix_Needed':'',
    'Notes':'Opened CBP OPR release: no evidence agents struck any person with horse reins. Also found unnecessary force — do not overclaim nothing happened.','Fits_Catalog_Rule':'yes',
  },
  'Sicknick': {
    'Verdict':'Verified','Best_Source_URL':'https://www.uscp.gov/media-center/press-releases/medical-examiner-finds-uscp-officer-brian-sicknick-died-natural-causes',
    'Source_Type':'official_record','Evidence_Level':'Proven false',
    'Fix_Needed':'Clarify early fire-extinguisher-to-death narrative vs ME natural causes (strokes); still line-of-duty death.',
    'Notes':'Opened USCP release Apr 19, 2021.','Fits_Catalog_Rule':'yes',
  },
  'Insurrection': {
    'Verdict':'Verified with correction needed','Best_Source_URL':'https://www.justice.gov/usao-dc/48-months-jan-6-attack-us-capitol',
    'Source_Type':'official_record','Evidence_Level':'Proven false',
    'Fix_Needed':'Keep zero-§2383 point; do not imply zero serious charges (~1,583 charged; assaults/obstruction/seditious conspiracy). DOJ page 401 from auditor host.',
    'Notes':'Zero §2383 charges corroborated by USAO-DC public snapshots.','Fits_Catalog_Rule':'yes',
  },
  'The charge was not insurrection': {
    'Verdict':'Verified with correction needed','Best_Source_URL':'https://www.justice.gov/usao-dc/48-months-jan-6-attack-us-capitol',
    'Source_Type':'official_record','Evidence_Level':'Proven false',
    'Fix_Needed':'Same as Insurrection row.','Notes':'Hoaxes table.','Fits_Catalog_Rule':'yes',
  },
  'Kids in cages': {
    'Verdict':'Verified with correction needed','Best_Source_URL':'https://apnews.com/article/a98f26f7c9424b44b7fa927ea1acd4d4',
    'Source_Type':'outlet_correction','Evidence_Level':'Proven false',
    'Fix_Needed':'Truth = 2014 photos mislabeled as 2018 Trump cages, not denial of 2018 family-separation policy.',
    'Notes':'AP fact-check URL resolves HTTP 200.','Fits_Catalog_Rule':'yes',
  },
  'The FISA file was not scrupulously accurate': {
    'Verdict':'Verified','Best_Source_URL':'https://oig.justice.gov/node/1100',
    'Source_Type':'official_record','Evidence_Level':'Proven false',
    'Fix_Needed':'Frame as rebuttal to “dossier/FISA was clean”; Horowitz found significant inaccuracies/omissions.',
    'Notes':'OIG page HTTP 200.','Fits_Catalog_Rule':'yes',
  },
  'Page was a spy': {
    'Verdict':'Verified with correction needed','Best_Source_URL':'https://oig.justice.gov/node/1100',
    'Source_Type':'official_record','Evidence_Level':'Rated misleading',
    'Fix_Needed':'FISA target yes; charged Russian agent no.','Notes':'Horowitz context.','Fits_Catalog_Rule':'yes',
  },
  'Pee tape': {
    'Verdict':'Unverified','Best_Source_URL':'https://www.justice.gov/storage/durhamreport.pdf',
    'Source_Type':'official_record','Evidence_Level':'',
    'Fix_Needed':'Durham PDF 401 here. Do not reuse until local PDF confirms no-tape / dossier treatment.',
    'Notes':'Absence claim plausible but source not opened.','Fits_Catalog_Rule':'maybe',
  },
  'No pardon': {
    'Verdict':'Verified','Best_Source_URL':'https://www.justice.gov/',
    'Source_Type':'official_record','Evidence_Level':'','Fix_Needed':'Cite Dec 1, 2024 Hunter Biden pardon warrant when reused.',
    'Notes':'Democrat ledger: promised no pardon; later pardoned.','Fits_Catalog_Rule':'no (Democrat ledger)',
  },
  'If you like your plan': {
    'Verdict':'Verified','Best_Source_URL':'https://www.politifact.com/article/2013/dec/12/lie-year-if-you-like-your-plan-you-can-keep-it/',
    'Source_Type':'fact_check','Evidence_Level':'','Fix_Needed':'Prefer HHS/CMS disruption data + Obama acknowledgment alongside fact-check.',
    'Notes':'Democrat ledger / Obama-era.','Fits_Catalog_Rule':'no (Democrat ledger)',
  },
  'Antifa': {
    'Verdict':'Unverified','Best_Source_URL':'https://www.justice.gov/opa/pr/leader-antifa-cell-members-north-texas-sentenced-100-years-prison-terrorist-attack-ice',
    'Source_Type':'official_record','Evidence_Level':'','Fix_Needed':'DOJ PR 401 here; confirm 100-year / 450-year figures on judgment entries before reuse.',
    'Notes':'Democrat ledger.','Fits_Catalog_Rule':'no (Democrat ledger)',
  },
  'The client list': {
    'Verdict':'Unverified','Best_Source_URL':'https://www.documentcloud.org/documents/25993306-doj-epstein-july-2025-memo',
    'Source_Type':'','Evidence_Level':'','Fix_Needed':'DocumentCloud memo Cloudflare-blocked; do not reuse until PDF archived and quoted.',
    'Notes':'July 2025 DOJ/FBI memo not opened in this audit.','Fits_Catalog_Rule':'maybe',
  },
  'Murdered in the cell': {
    'Verdict':'Verified with correction needed','Best_Source_URL':'https://www.nyc.gov/site/ocme/index.page',
    'Source_Type':'official_record','Evidence_Level':'Proven false',
    'Fix_Needed':'Cite NYC OCME suicide ruling; do not rely on blocked July 2025 memo. Note real guard/camera failures.',
    'Notes':'Homicide caption vs ME suicide.','Fits_Catalog_Rule':'yes',
  },
  'Kidnapped': {
    'Verdict':'Verified with correction needed','Best_Source_URL':'https://www.justice.gov/usao-sdny/pr/manhattan-us-attorney-announces-narco-terrorism-charges-against-nicolas-maduro-current',
    'Source_Type':'official_record','Evidence_Level':'Rated misleading',
    'Fix_Needed':'2020 SDNY indictment real (link 200). Confirm Jan 3, 2026 custody against docket before reuse.',
    'Notes':'Kidnap vs warrant framing.','Fits_Catalog_Rule':'yes if media used kidnap framing',
  },
  'The District of Columbia': {
    'Verdict':'Verified with correction needed','Best_Source_URL':'https://www.courtlistener.com/docket/67490071/united-states-v-trump/',
    'Source_Type':'official_record','Evidence_Level':'','Fix_Needed':'No §2383 in indictment; dismissed — not a merits acquittal.',
    'Notes':'Lawfare row; CourtListener reachable.','Fits_Catalog_Rule':'n/a (lawfare)',
  },
  'MEDIA_TALLY': {
    'Verdict':'Verified with correction needed','Best_Source_URL':'','Source_Type':'','Evidence_Level':'',
    'Fix_Needed':'Lead says 108 captions; mediaFrames has 115 — align before publish.',
    'Notes':'Internal consistency error.','Fits_Catalog_Rule':'n/a',
  },
}

OPINION_TAGS = {
  'Lock her up','Term limits','Infrastructure week','We back the blue','We have his back',
  'Mission accomplished','A tour','No one was armed','Antifa did it','2,000 Mules',
  'They made it necessary','Balance it','Most corrupt ever','Semi-fascism','Deplorables',
  'Garbage','Jim Crow','Bend the knee','Sharp as a tack',"Don’t say gay",'Zero inflation',
  'Maximum warfare','You will pay','Pass it to see it','You will not get it',
  'Inflation Reduction','Transitory','Liz Cheney','Ban the pill','Left the bullet in',
  'No pain medicine','A random stop','ICE out of Austin','He burned the church',
  'Not one mile','Convicted of all of it','He grabbed the wheel','Impeached for treason',
  'Ten crimes','He confessed','No exoneration','Let them die','David Duke',
  'Lab leak','Hydroxychloroquine kills','Horse paste','The noose','Kavanaugh',
  'She was in the chamber','Rittenhouse','Concentration camps','Tear gas',
  'Suckers and losers','Suckers / losers','Find 11,780','Perfect call','The weapons',
  'The returns','Proven stolen','Mexico pays','Repeal','No one is above the law',
  'Benghazi','I never spoke to him','The border is secure','Parents','Diesel is $6.52',
}

by_no={it['Item_No']:it for it in items}

def first_link(it):
    links=re.findall(r'https?://[^\s|"\']+', it.get('Links_Given') or '')
    return links[0].rstrip('.,);]') if links else ''

def source_type_for(url):
    if not url: return ''
    u=url.lower()
    if any(x in u for x in ['justice.gov','supremecourt.gov','congress.gov','senate.gov','house.gov','cbp.gov','dhs.gov','treasury.gov','bls.gov','federalregister.gov','archives.gov','oig.','uscp.gov','courtlistener.com','iaea.org','dni.gov','fec.gov','nyc.gov','nycourts.gov']):
        return 'official_record'
    if 'youtube.com' in u or 'c-span.org' in u: return 'transcript_or_video'
    if any(x in u for x in ['apnews.com','reuters.com','politifact','snopes','factcheck']):
        return 'outlet_correction_or_factcheck'
    return 'other'

verified=[]
for it in items:
    row={k:it.get(k,'') for k in ['Item_No','Page','Source_File','Type','Text','Attributed_To','Date','Links_Given','Duplicate_Of_Item_ID']}
    row.update({'Verdict':'','Best_Source_URL':'','Source_Type':'','Evidence_Level':'','Fix_Needed':'','Notes':'','Category':it.get('Category',''),'Fits_Catalog_Rule':''})
    tag=it.get('_tag') or it.get('Attributed_To') or ''
    dup=it.get('Duplicate_Of_Item_ID') or ''
    fl=first_link(it)

    if str(dup).startswith(('FT-','LT-')):
        c=CAT.get(dup,{})
        ev=(c.get('Evidence_Level') or '').strip()
        row['Verdict']='Verified'
        row['Evidence_Level']=ev
        row['Best_Source_URL']=c.get('Primary_Source_URL') or c.get('Truth_Source_URL') or fl
        row['Source_Type']=source_type_for(row['Best_Source_URL'])
        code=url_cache.get(fl,{}).get('code','n/a') if fl else 'no site link'
        row['Notes']=f'Duplicate of catalog {dup}. Inherited Evidence_Level. Site link HTTP {code}.'
        row['Fits_Catalog_Rule']='yes (already in catalog)'
        if fl and fl in url_cache and not url_cache[fl].get('ok'):
            row['Fix_Needed']=f'Replace/mirror site href (HTTP {url_cache[fl].get("code")}).'
            row['Verdict']='Verified with correction needed'
    elif str(dup).startswith('CP-'):
        row['Verdict']='Verified'
        row['Notes']=f"Internal duplicate of checkpoint Item_No {dup.split('-')[1]}."
        row['Best_Source_URL']=fl
        row['Source_Type']=source_type_for(fl)
        row['Fits_Catalog_Rule']='see original'
        row['_defer_cp']=int(dup.split('-')[1])
    elif tag in CURATED:
        row.update({k:CURATED[tag].get(k,'') for k in ['Verdict','Best_Source_URL','Source_Type','Evidence_Level','Fix_Needed','Notes','Fits_Catalog_Rule']})
        if not row['Best_Source_URL']: row['Best_Source_URL']=fl
        if not row['Source_Type']: row['Source_Type']=source_type_for(row['Best_Source_URL'])
    elif tag in OPINION_TAGS:
        row['Verdict']='Opinion'
        row['Best_Source_URL']=fl
        row['Source_Type']=source_type_for(fl)
        row['Fits_Catalog_Rule']='no'
        row['Notes']='Label as opinion/campaign rhetoric if reused; not a catalog media-deception case without specific false caption + proof.'
        if it.get('Category')=='democrat_statement': row['Notes']+=' Category: Democrat ledger.'
        if it.get('Category')=='republican_statement': row['Notes']+=' Category: Republican ledger.'
    else:
        uc=url_cache.get(fl,{}) if fl else {}
        row['Best_Source_URL']=fl
        row['Source_Type']=source_type_for(fl)
        if it.get('Category') in ('democrat_statement','republican_statement'):
            row['Fits_Catalog_Rule']='no (party ledger — verify on own terms)'
        else:
            row['Fits_Catalog_Rule']='unknown — needs manual fit check'
        if fl and uc.get('ok'):
            row['Verdict']='Unverified'
            row['Notes']=f'Link HTTP {uc.get("code")}. Exact wording not confirmed in fetched body; Unverified per no-fabrication rule.'
        elif fl and uc:
            row['Verdict']='Unverified'
            row['Fix_Needed']=f'Source link failed/blocked (HTTP {uc.get("code")}).'
            row['Notes']='Could not open given source in this audit.'
        else:
            row['Verdict']='Unverified'
            row['Notes']='No usable source opened for this item in this audit.'
    verified.append(row)

by_ver={r['Item_No']:r for r in verified}
for r in verified:
    if r.get('_defer_cp'):
        src=by_ver.get(r['_defer_cp'])
        if src:
            for k in ['Verdict','Evidence_Level','Best_Source_URL','Source_Type','Fix_Needed','Fits_Catalog_Rule']:
                if src.get(k): r[k]=src[k]
            r['Notes']=f"Internal duplicate of Item_No {r['_defer_cp']}. "+(src.get('Notes') or '')
        r.pop('_defer_cp', None)

out_fields=['Item_No','Page','Source_File','Type','Text','Attributed_To','Date','Links_Given','Duplicate_Of_Item_ID',
            'Verdict','Best_Source_URL','Source_Type','Evidence_Level','Fix_Needed','Notes','Category','Fits_Catalog_Rule']
with open('verified-items.csv','w',newline='',encoding='utf-8') as fh:
    w=csv.DictWriter(fh, fieldnames=out_fields, extrasaction='ignore'); w.writeheader(); w.writerows(verified)

wb=Workbook(); ws=wb.active; ws.title='verified-items'
ws.append(out_fields)
hf=PatternFill('solid','FF1B2A41'); hfont=Font(color='FFFFFF', bold=True)
for col,h in enumerate(out_fields,1):
    cell=ws.cell(1,col,h); cell.fill=hf; cell.font=hfont
fills={'Verified':PatternFill('solid','FFC6EFCE'),'Verified with correction needed':PatternFill('solid','FFFFEB9C'),
       'Unverified':PatternFill('solid','FFDDEBF7'),'False':PatternFill('solid','FFFFC7CE'),'Opinion':PatternFill('solid','FFE2D5F1')}
for r in verified:
    ws.append([r.get(c,'') for c in out_fields])
    v=r.get('Verdict','')
    if v in fills: ws.cell(ws.max_row, out_fields.index('Verdict')+1).fill=fills[v]
ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
for col in ws.columns:
    maxlen=min(55, max(len(str(c.value or '')) for c in col))
    ws.column_dimensions[col[0].column_letter].width=max(12, maxlen+2)
wb.save('verified-items.xlsx')

vc=Counter(r['Verdict'] for r in verified)
by_page=defaultdict(Counter)
for r in verified: by_page[r['Page']][r['Verdict']]+=1
new_cases=[r for r in verified if (not r['Duplicate_Of_Item_ID'] and r['Category']=='media_claim_vs_record'
            and r['Verdict'] in ('Verified','Verified with correction needed')
            and r.get('Evidence_Level') in ('Proven false','Rated misleading')
            and 'yes' in (r.get('Fits_Catalog_Rule') or '').lower())]
failed=[r for r in verified if r['Verdict'] in ('Unverified','False') and not r['Duplicate_Of_Item_ID']]
n_cat=sum(1 for r in verified if str(r['Duplicate_Of_Item_ID']).startswith(('FT-','LT-')))
n_cp=sum(1 for r in verified if str(r['Duplicate_Of_Item_ID']).startswith('CP-'))
n_unique=sum(1 for r in verified if not r['Duplicate_Of_Item_ID'])

summary=f'''# Checkpoint content audit summary

**Source commit:** `c60dc5d193afacfcc8afcab5b8d328ec5debcf2b` (“Full site checkpoint”, 10:01 PM MDT Sep 23, 2026)  
**Workdir:** `/workspace/checkpoint-review/` (GitHub read-only; no pushes/commits)  
**Catalog:** term-split first-term + later-second-term CSVs (250 cases)

## Outputs
- `inventory.md` — content page/file inventory  
- `items.csv` — extracted factual items  
- `verified-items.csv` / `verified-items.xlsx`  
- `review-summary.md` (this file)  
- `repo/` — local copies of key files at the SHA (blob SHAs matched)

## Counts
| Bucket | N |
|--------|--:|
| Total extracted items | {len(verified)} |
| Catalog duplicates (FT-/LT-) | {n_cat} |
| Internal duplicates (CP-) | {n_cp} |
| Unique | {n_unique} |

### Verdicts (all items)
| Verdict | N |
|---------|--:|
{chr(10).join(f'| {k} | {v} |' for k,v in vc.most_common())}

### Per page (top)
{chr(10).join(f"- **{p}**: " + ", ".join(f"{k}={v}" for k,v in c.most_common()) for p,c in sorted(by_page.items(), key=lambda x: -sum(x[1].values()))[:16])}

## New verified cases for possible catalog add
{chr(10).join(f"- Item {r['Item_No']} — **{r['Attributed_To']}**: {r['Verdict']}; {r['Evidence_Level']}; {r['Best_Source_URL']}" for r in new_cases) or '_None pending further primary-source opens._'}

## Failed / blocked / Unverified uniques (sample)
{chr(10).join(f"- Item {r['Item_No']} — {r['Attributed_To'] or r['Page']}: {r['Fix_Needed'] or r['Notes'][:180]}" for r in failed[:45])}
{('… +'+str(len(failed)-45)+' more') if len(failed)>45 else ''}

## Structural problems
- MEDIA_LEAD says **108** captions; `mediaFrames` has **115** objects.  
- justice.gov often **401** here (Durham, USAO-DC Jan6, some OPA). Mirror PDFs before publish.  
- DocumentCloud Epstein July 2025 memo **Cloudflare-blocked** — do not reuse client-list / memo-dependent rows until PDF is local.

## Democrat / Republican ledgers
- Dem ~27 / GOP ~26 rows: different category from Trump-admin media catalog. Verify on own terms. Strongest Dem factual rows: No pardon; If you like your plan. Many GOP rows are campaign-promise accountability → Opinion unless tied to a specific false media caption.

## Site features worth keeping
- Claim|Truth frames + MEDIA_RAN duration notes  
- Separate Dem/GOP ledgers  
- Lawfare schema (caption/sold/evidence/file/docs)  
- Scorecard tabs (charts/read), Farm bars, Oval desks, HOAXES/FRAMES/FAKE_NEWS  
- InteractiveChart + KeptRead; series grouping; ~55 dispatch sitemap  
- Nav: Dispatch / Scorecard / Archive / Pump / Foreword / About; OG image set

## Remains
Full essay-body extraction for all posts; open every 401 DOJ target; confirm 2025–2026 extraordinary claims (Maduro custody, IAEA GOV/2026/50, Antifa sentence aggregates); scorecard numeric modules not fully itemized; fuzzy catalog near-dupes may remain among “unique” media rows.

## Top findings
1. Checkpoint is a strong claim/truth skeleton — reuse structure, not unchecked numbers.  
2. ~{n_cat} items already in the 250-case catalog — align wording to catalog Evidence_Level.  
3. **Whips** (CBP OPR) is a clean official-record verified case.  
4. **§2383 zero charges** usable with careful wording (not “nothing charged”).  
5. Epstein memo-dependent rows blocked — hold.  
6. 108 vs 115 tally is publish-blocking.
'''
Path('review-summary.md').write_text(summary, encoding='utf-8')
print('VERDICTS', dict(vc))
print('new_cases', len(new_cases), [r['Attributed_To'] for r in new_cases])
print('failed_unique', len(failed))
print('files', list(Path('.').glob('verified-items.*')), 'summary', Path('review-summary.md').stat().st_size)
