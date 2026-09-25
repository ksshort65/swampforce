#!/usr/bin/env python3
"""Weekly update for gas-gap-tracker.xlsx.
Re-downloads the free EIA weekly series and the CFTC current-year disaggregated COT file, then APPENDS only weeks
newer than the last row of each tab. Existing rows are never rewritten. Formula columns for new rows are copied
from the previous row (row references shifted). A dated backup of the workbook is written first.
Run:  /workspace/energy/.venv/bin/python /workspace/energy/gap-tracker/update.py
"""
import io, sys, shutil, zipfile, datetime as dt, urllib.request
import pandas as pd
from openpyxl import load_workbook
from openpyxl.formula.translate import Translator
G = '/workspace/energy/gap-tracker/'; XL = G + 'gas-gap-tracker.xlsx'
UA = {'User-Agent': 'Mozilla/5.0 (SwampForce gap tracker; editor@swampforce.com)'}
FUELTAX = 52.16  # EIA avg federal+state gasoline tax, Jul 2026 (update by hand when EIA publishes a new fueltaxes.xlsx)
def get(url):
    return urllib.request.urlopen(urllib.request.Request(url, headers=UA), timeout=120).read()
def eia(series):
    b = get(f'https://www.eia.gov/dnav/pet/hist_xls/{series}w.xls')
    x = pd.read_excel(io.BytesIO(b), sheet_name='Data 1', skiprows=2); x.columns = ['date', 'v']; x = x.dropna()
    x['date'] = pd.to_datetime(x.date); return x.set_index('date').v
def last_date(ws, first_row):
    d = None
    for r in range(first_row, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if isinstance(v, (dt.date, dt.datetime)): d = v
    return pd.Timestamp(d), max(r for r in range(first_row, ws.max_row + 1) if isinstance(ws.cell(r, 1).value, (dt.date, dt.datetime)))
def append(ws, last_row, values, n_raw, fmt_cols=(1,)):
    r = last_row + 1
    for j, v in enumerate(values, 1):
        if v is not None and not (isinstance(v, float) and pd.isna(v)): ws.cell(r, j, v)
    for j in fmt_cols: ws.cell(r, j).number_format = 'yyyy-mm-dd'
    for j in range(n_raw + 1, ws.max_column + 1):
        f = ws.cell(last_row, j).value
        if isinstance(f, str) and f.startswith('='):
            ws.cell(r, j, Translator(f, origin=ws.cell(last_row, j).coordinate).translate_formula(ws.cell(r, j).coordinate))
    return r
def main():
    wb = load_workbook(XL); log = []
    # Weekly
    S = ['RWTC', 'RBRTE', 'EER_EPMRU_PF4_RGC_DPG', 'EER_EPMRU_PF4_Y35NY_DPG', 'EER_EPMRR_PF4_Y05LA_DPG', 'EMM_EPMR_PTE_NUS_DPG', 'EMM_EPMR_PTE_SCO_DPG', 'EMM_EPMR_PTE_R40_DPG', 'EMM_EPMR_PTE_SCA_DPG']
    D = {s: eia(s) for s in S}; retail = set(S[5:])
    ws = wb['Weekly 2005-present']; ld, lr = last_date(ws, 2)
    for d in D['EMM_EPMR_PTE_NUS_DPG'][D['EMM_EPMR_PTE_NUS_DPG'].index > ld].index:
        f = d - pd.Timedelta(days=3)
        vals = [d.date(), f.date()] + [(float(D[s][d]) if d in D[s].index else None) if s in retail else (float(D[s][f]) if f in D[s].index else None) for s in S]
        vals += [None, FUELTAX]  # EIA monthly tax share not yet published for new weeks -> use EIA avg tax
        lr = append(ws, lr, vals, len(vals), fmt_cols=(1, 2)); log.append(f'Weekly +{d.date()}')
    # Exports vs inventories
    E = ['WGTSTUS1', 'W_EPM0F_EEX_NUS-Z00_MBBLD', 'WRPEXUS2', 'WCREXUS2']; ED = {s: eia(s) for s in E}
    ws = wb['Exports vs inventories']; ld, lr = last_date(ws, 2)
    for d in ED['WGTSTUS1'][ED['WGTSTUS1'].index > ld].index:
        vals = [d.date()] + [float(ED[s][d]) if d in ED[s].index else None for s in E]
        lr = append(ws, lr, vals, 5); log.append(f'Exports +{d.date()}')
    # CFTC (current year file)
    y = dt.date.today().year
    z = zipfile.ZipFile(io.BytesIO(get(f'https://www.cftc.gov/files/dea/history/fut_disagg_txt_{y}.zip')))
    c = pd.read_csv(z.open(z.namelist()[0]), low_memory=False)
    c['code'] = c.CFTC_Contract_Market_Code.astype(str).str.strip(); c = c[c.code.isin(['067651', '111659'])].copy()
    c['date'] = pd.to_datetime(c['Report_Date_as_YYYY-MM-DD']); c = c.sort_values(['date', 'code'])
    ws = wb['Futures positioning (CFTC)']; ld, lr = last_date(ws, 3)
    for r in c[c.date > ld].itertuples():
        vals = [r.date.date(), r.code, r.Market_and_Exchange_Names.strip(), int(r.Open_Interest_All), int(r.M_Money_Positions_Long_All), int(r.M_Money_Positions_Short_All), int(r.M_Money_Positions_Spread_All)]
        lr = append(ws, lr, vals, 7); log.append(f'CFTC +{r.date.date()} {r.code}')
    if not log:
        print('No new weeks; workbook unchanged.'); return 0
    shutil.copy(XL, G + f'gas-gap-tracker.backup-{dt.date.today()}.xlsx')
    wb.save(XL)
    with open(G + 'update-log.txt', 'a') as fh:
        fh.write(f'{dt.datetime.now():%Y-%m-%d %H:%M} appended {len(log)} rows: ' + ', '.join(log) + '\n')
    print(f'Appended {len(log)} rows.', *log, sep='\n  ')
if __name__ == '__main__':
    sys.exit(main())
