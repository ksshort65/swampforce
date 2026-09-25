"""Builds gas-gap-tracker.xlsx, charts and facts.json from files already in /workspace/energy/raw (no downloads)."""
import pandas as pd, numpy as np, json, glob, collections, datetime as dt
import matplotlib; matplotlib.use('Agg'); import matplotlib.pyplot as plt, matplotlib.dates as mdates
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as L
R='/workspace/energy/raw/'; G='/workspace/energy/gap-tracker/'
DN='https://www.eia.gov/dnav/pet/hist/LeafHandler.ashx?n=PET&s={}&f={}'
def rd(f):
    x=pd.read_excel(R+f+'.xls',sheet_name='Data 1',skiprows=2); x.columns=['date','v']; x=x.dropna(); x['date']=pd.to_datetime(x.date); return x.set_index('date').v
F={}
# ---------------- weekly
W=dict(wti=rd('RWTCw'),brent=rd('RBRTEw'),gc=rd('EER_EPMRU_PF4_RGC_DPGw'),ny=rd('EER_EPMRU_PF4_Y35NY_DPGw'),la=rd('EER_EPMRR_PF4_Y05LA_DPGw'),
       us=rd('EMM_EPMR_PTE_NUS_DPGw'),co=rd('EMM_EPMR_PTE_SCO_DPGw'),p4=rd('EMM_EPMR_PTE_R40_DPGw'),ca=rd('EMM_EPMR_PTE_SCA_DPGw'))
pc=pd.read_csv(R+'eia_gas_pump_components_history_raw.csv')
mm={'Jan':1,'Feb':2,'Mar':3,'Apr':4,'May':5,'Jun':6,'Jul':7,'July':7,'Aug':8,'Sep':9,'Sept':9,'Oct':10,'Nov':11,'Dec':12}
pc['month']=pc.mon_yr.apply(lambda s: pd.Timestamp(2000+int(s.split('-')[1]),mm[s.split('-')[0]],1)); pc=pc.sort_values('month').reset_index(drop=True)
taxshare={r.month:r.taxes_pct for r in pc.itertuples()}; last_pc=pc.month.max()
FUELTAX=52.16  # EIA fueltaxes.xlsx, avg federal+state, Jul 2026
wk=W['us'][W['us'].index>='2005-01-03']
SER=[('WTI Cushing spot, $/bbl (EIA RWTC weekly, week ending Fri)','wti','RWTC'),('Brent spot, $/bbl (EIA RBRTE weekly)','brent','RBRTE'),
     ('Gulf Coast conv. regular gasoline spot, $/gal (EIA)','gc','EER_EPMRU_PF4_RGC_DPG'),('NY Harbor conv. regular gasoline spot, $/gal (EIA)','ny','EER_EPMRU_PF4_Y35NY_DPG'),
     ('Los Angeles RBOB regular spot, $/gal (EIA)','la','EER_EPMRR_PF4_Y05LA_DPG'),('U.S. retail regular, $/gal (EIA, Monday)','us','EMM_EPMR_PTE_NUS_DPG'),
     ('Colorado retail regular, $/gal (EIA)','co','EMM_EPMR_PTE_SCO_DPG'),('Rocky Mountain PADD 4 retail regular, $/gal (EIA)','p4','EMM_EPMR_PTE_R40_DPG'),('California retail regular, $/gal (EIA)','ca','EMM_EPMR_PTE_SCA_DPG')]
rows=[]
for d,v in wk.items():
    f=d-pd.Timedelta(days=3); m=pd.Timestamp(d.year,d.month,1)
    r=[d.date(),f.date()]+[(float(W[k].get(f)) if k not in('us','co','p4','ca') else float(W[k].get(d))) if (W[k].get(f if k not in('us','co','p4','ca') else d) is not None) else None for _,k,_ in SER]
    r+=[float(taxshare[m]) if m in taxshare else None, None if m in taxshare else FUELTAX]
    rows.append(r)
wdf=pd.DataFrame(rows,columns=['week','spot_week']+[k for _,k,_ in SER]+['taxpct','fueltax'])
wdf['tax']=np.where(wdf.taxpct.notna(),wdf.us*wdf.taxpct/100,wdf.fueltax/100)
wdf['gap']=wdf.us-wdf.wti/42-wdf.tax; wdf['crack']=wdf.gc-wdf.wti/42; wdf['crackb']=wdf.gc-wdf.brent/42; wdf['ret_whl']=wdf.us-wdf.gc
# ---------------- workbook
wb=Workbook(); HF=Font(bold=True,color='FFFFFF'); HP=PatternFill('solid',fgColor='1E3A5F'); FP=PatternFill('solid',fgColor='FFF2CC')
def sheet(name,headers,data,formulas=(),widths=None,note=None):
    ws=wb.create_sheet(name); r0=1
    if note: ws.cell(1,1,note).font=Font(italic=True); r0=2
    allh=list(headers)+[h for h,_ in formulas]
    for j,h in enumerate(allh,1):
        c=ws.cell(r0,j,h); c.font=HF; c.fill=FP if h.endswith('[formula]') else HP; c.alignment=Alignment(wrap_text=True,vertical='top')
        if h.endswith('[formula]'): c.font=Font(bold=True)
    for i,row in enumerate(data):
        rr=r0+1+i
        for j,v in enumerate(row,1):
            if v is not None and not (isinstance(v,float) and np.isnan(v)): ws.cell(rr,j,v)
        for k,(h,fn) in enumerate(formulas):
            ws.cell(rr,len(headers)+1+k,fn(rr))
    for j in range(1,len(allh)+1): ws.column_dimensions[L(j)].width=(widths[j-1] if widths and j-1<len(widths) else 16)
    ws.row_dimensions[r0].height=60; ws.freeze_panes=ws.cell(r0+1,1); return ws
wb.remove(wb.active)
readme=wb.create_sheet('README')
for i,t in enumerate(["Gas Price Gap Tracker (SwampForce). Built Sep 24, 2026 (MT) from primary records.",
 "Raw columns (dark header) are typed exactly as reported by the source and are never altered.",
 "Yellow columns ending in [formula] are Excel formulas computed from raw columns in the same row. Nothing else is derived.",
 "Every series or row links to its source on the 'Sources' tab or in its own URL column.",
 "The 'Unexplained - Red flags' tab (Excel does not allow '/' in tab names) holds fixed text/values summarizing data points that do not fit the documented explanations. A red flag is not a finding of wrongdoing.",
 "Weekly: retail prices are EIA Monday readings; spot prices are EIA weekly averages for the week ending the prior Friday.",
 "Weekly taxes: EIA pump-components tax share (%) for the month (Jan 2000-May 2026); after May 2026 EIA's Jul 2026 average federal+state gasoline tax (52.16 c/gal) is used.",
 "update.py re-downloads the free EIA/CFTC series and appends only new weeks; existing rows are never rewritten."],1):
    readme.cell(i,1,t)
readme.column_dimensions['A'].width=160
hdr=['Retail week (Mon)','Spot week ending (Fri)']+[h for h,_,_ in SER]+['EIA pump-components tax share of retail, % (monthly; blank after May 2026)','EIA avg federal+state gasoline tax, c/gal (Jul 2026; used after May 2026)']
col={k:L(3+i) for i,(_,k,_) in enumerate(SER)}; TP=L(3+len(SER)); FT=L(4+len(SER))
fw=[('Gulf Coast gasoline minus WTI/42, $/gal (crack spread) [formula]',lambda r:f'=IF(OR({col["gc"]}{r}="",{col["wti"]}{r}=""),"",{col["gc"]}{r}-{col["wti"]}{r}/42)'),
    ('Gulf Coast gasoline minus Brent/42, $/gal [formula]',lambda r:f'=IF(OR({col["gc"]}{r}="",{col["brent"]}{r}=""),"",{col["gc"]}{r}-{col["brent"]}{r}/42)'),
    ('Taxes, $/gal [formula]',lambda r:f'=IF({TP}{r}="",{FT}{r}/100,{col["us"]}{r}*{TP}{r}/100)'),
    ('GAP = U.S. retail - WTI/42 - taxes, $/gal [formula]',lambda r:f'={col["us"]}{r}-{col["wti"]}{r}/42-{L(3+len(SER)+4)}{r}'),
    ('U.S. retail minus Gulf Coast spot, $/gal [formula]',lambda r:f'={col["us"]}{r}-{col["gc"]}{r}'),
    ('Colorado minus U.S. retail, $/gal [formula]',lambda r:f'=IF({col["co"]}{r}="","",{col["co"]}{r}-{col["us"]}{r})')]
data=[list(r[:2])+list(r[2:]) for r in wdf[['week','spot_week']+[k for _,k,_ in SER]+['taxpct','fueltax']].itertuples(index=False)]
ws=sheet('Weekly 2005-present',hdr,data,fw)
for row in ws.iter_rows(min_row=2,max_col=2):
    for c in row: c.number_format='yyyy-mm-dd'
# monthly components
cdata=[[r.month.date(),r.mon_yr,r.retail_usd_gal,r.crude_pct,r.refining_pct,r.dist_marketing_pct,r.taxes_pct] for r in pc.itertuples()]
ws=sheet('Monthly pump components',['Month','EIA label','Retail regular, $/gal (EIA)','Crude oil, % of retail (EIA)','Refining costs & profits, % (EIA)','Distribution & marketing, % (EIA)','Taxes, % (EIA)'],cdata,
   [('Crude, c/gal [formula]',lambda r:f'=C{r}*D{r}'),('Refining, c/gal [formula]',lambda r:f'=C{r}*E{r}'),('Distribution & marketing, c/gal [formula]',lambda r:f'=C{r}*F{r}'),('Taxes, c/gal [formula]',lambda r:f'=C{r}*G{r}'),('Shares sum, % [formula]',lambda r:f'=D{r}+E{r}+F{r}+G{r}')])
for row in ws.iter_rows(min_row=2,max_col=1):
    for c in row: c.number_format='yyyy-mm'
# ---------------- refiners (SEC XBRL)
CO=[('Valero Energy','0001035002','us-gaap','USD',None),('Marathon Petroleum','0001510295','us-gaap','USD',None),('Phillips 66','0001534701','us-gaap','USD',None),
    ('PBF Energy','0001534504','us-gaap','USD',None),('HollyFrontier (to 2021)','0000048039','us-gaap','USD',(2008,2021)),('HF Sinclair (2022-)','0001915657','us-gaap','USD',(2022,2026)),('Suncor Energy (whole company, CAD, IFRS; 40-F)','0000311337','ifrs-full','CAD',None)]
TAGS={'us-gaap':{'Net income attributable to company':['NetIncomeLoss'],'Share buybacks (cash paid)':['PaymentsForRepurchaseOfCommonStock','PaymentsForRepurchaseOfEquity'],'Dividends paid to common (cash)':['PaymentsOfDividendsCommonStock','PaymentsOfDividends']},
      'ifrs-full':{'Net income (profit or loss, incl. minority)':['ProfitLoss'],'Share buybacks (cash paid)':['PaymentsToAcquireOrRedeemEntitysShares'],'Dividends paid (cash)':['DividendsPaidClassifiedAsFinancingActivities']}}
ref=[]
for name,cik,tax,cur,yr in CO:
    fct=json.load(open(R+f'cf_{cik}.json'))['facts'].get(tax,{})
    for metric,tags in TAGS[tax].items():
        got={}
        for t in tags:
            if t not in fct: continue
            for u in fct[t]['units'].get(cur,[]):
                per=None
                if u.get('frame','').startswith('CY') and len(u['frame'])==6: per=u['frame']
                elif u.get('frame') in ('CY2026Q2','CY2025Q2') and metric.startswith('Net'): per=u['frame']
                elif u.get('start')=='2026-01-01' and u.get('end')=='2026-06-30' and u.get('form','').startswith('10-Q'): per='2026 H1 (Jan-Jun)'
                if per and per not in got:
                    y=int(per[2:6]) if per.startswith('CY') else 2026
                    if yr and not(yr[0]<=y<=yr[1]): continue
                    if y<2008 or (per.startswith('CY') and len(per)==6 and y not in (2008,)+tuple(range(2012,2026))): continue
                    got[per]=(u['val'],u.get('form'),u.get('accn'),t)
        for per,(v,form,accn,t) in sorted(got.items()):
            url=f"https://www.sec.gov/Archives/edgar/data/{int(cik)}/{accn.replace('-','')}/"
            ref.append([name,metric,per.replace('CY','') if per.startswith('CY') and len(per)==6 else per.replace('CY',''),v,cur,form,accn,t,url])
rdf=pd.DataFrame(ref,columns=['co','metric','period','val','cur','form','accn','tag','url'])
ws=sheet('Refiners (SEC)',['Company','Metric','Period (calendar year unless noted)','Value as reported (units of currency)','Currency','SEC form','Accession no.','XBRL tag','Filing URL'],ref,
   [('Value, $ billions [formula]',lambda r:f'=D{r}/1e9')],widths=[26,34,18,20,8,8,22,34,60],
   note='Source: SEC EDGAR XBRL company facts (data.sec.gov/api/xbrl/companyfacts). Values exactly as tagged in each filing. Suncor files IFRS in CAD, company-wide (oil sands + refining); Commerce City is not reported separately. PBF and HF Sinclair 2026 H1 as filed in Q2 10-Qs.')
# ---------------- ownership
cz=json.load(open(G+'refcap_concentration.json'))
T5='https://www.eia.gov/petroleum/refinerycapacity/table5.pdf'
own=[['United States',i+1,c,v,cz['national']['total'],T5] for i,(c,v) in enumerate(cz['national']['top'])]
for p in ['1','2','3','4','5']:
    d=cz['padd'][p]; own+= [[f'PADD {p}',i+1,c,v,d['total'],T5] for i,(c,v) in enumerate(d['top'])]
own.append(['Colorado',1,'SUNCOR ENERGY INC (Commerce City East 36,000 + West 67,000 b/cd)',103000,103000,T5])
ws=sheet('Ownership concentration',['Area','Rank','Corporation (EIA Table 5 parent)','Operable crude distillation capacity, b/cd (EIA, Jan 1 2026)','Area total capacity, b/cd (sum of EIA Table 5 plants in area; U.S. total = EIA U.S. Total 18,160,493)','Source'],own,
   [('Share of area capacity [formula]',lambda r:f'=D{r}/E{r}'),('Cumulative top-N share [formula]',lambda r:f'=SUMIFS(D:D,A:A,A{r},B:B,"<="&B{r})/E{r}')],widths=[14,6,48,20,24,50],
   note='EIA Refinery Capacity Report 2026, Table 5 (capacity by parent corporation and plant). PADD assignment by plant state (EIA PADD definitions). Colorado has one refinery complex (Suncor, Commerce City). Colorado pipeline supply: see Sources tab (secondary, labeled).')
ws2=wb['Ownership concentration']; base=ws2.max_row+2
ws2.cell(base,1,'PADD 4 (Rocky Mountain) plant list, EIA Table 5').font=Font(bold=True)
for i,(c,l,s,v) in enumerate(cz['padd']['4']['refineries']):
    for j,x in enumerate([c,l,s,v]): ws2.cell(base+1+i,j+1,x)
# ---------------- exports vs inventories
E={'stocks':rd('WGTSTUS1w'),'gasx':rd('W_EPM0F_EEX_NUS-Z00_MBBLDw'),'prodx':rd('WRPEXUS2w'),'crx':rd('WCREXUS2w')}
ed=pd.DataFrame(E); ed=ed[ed.index>='2005-01-01']
edata=[[d.date()]+[None if np.isnan(x) else float(x) for x in r] for d,r in zip(ed.index,ed.values)]
ws=sheet('Exports vs inventories',['Week ending (Fri)','U.S. total gasoline stocks, thousand bbl (EIA WGTSTUS1)','U.S. total motor gasoline exports, kb/d (EIA W_EPM0F_EEX_NUS-Z00_MBBLD; from Jun 2010)','U.S. total petroleum product exports, kb/d (EIA WRPEXUS2)','U.S. crude oil exports, kb/d (EIA WCREXUS2)'],edata,
   [('One week of gasoline exports (kb/d x 7) as share of gasoline stocks [formula]',lambda r:f'=IF(C{r}="","",C{r}*7/B{r})')])
for row in ws.iter_rows(min_row=2,max_col=1):
    for c in row: c.number_format='yyyy-mm-dd'
ann=[[int(d.year),float(v)] for d,v in rd('MGFEXUS2a').items() if d.year>=2005]
annp={int(d.year):float(v) for d,v in rd('MTPEXUS2a').items()}
ws.cell(1,8,'Annual (EIA MGFEXUS2 / MTPEXUS2)').font=Font(bold=True); ws.cell(2,8,'Year');ws.cell(2,9,'Finished motor gasoline exports, kb/d');ws.cell(2,10,'Finished petroleum products exports, kb/d')
for i,(y,v) in enumerate(ann): ws.cell(3+i,8,y); ws.cell(3+i,9,v); ws.cell(3+i,10,annp.get(y))
for c in 'HIJ': ws.column_dimensions[c].width=18
# ---------------- CFTC
fs=glob.glob(R+'cot/*/*'); cot=pd.concat([pd.read_csv(f,low_memory=False,usecols=['Market_and_Exchange_Names','Report_Date_as_YYYY-MM-DD','CFTC_Contract_Market_Code','Open_Interest_All','M_Money_Positions_Long_All','M_Money_Positions_Short_All','M_Money_Positions_Spread_All']) for f in fs])
cot['code']=cot.CFTC_Contract_Market_Code.astype(str).str.strip(); cot=cot[cot.code.isin(['067651','111659'])].copy()
cot['date']=pd.to_datetime(cot['Report_Date_as_YYYY-MM-DD']); cot=cot.drop_duplicates(['code','date']).sort_values(['code','date'])
cdat=[[r.date.date(),r.code,r.Market_and_Exchange_Names.strip(),int(r.Open_Interest_All),int(r.M_Money_Positions_Long_All),int(r.M_Money_Positions_Short_All),int(r.M_Money_Positions_Spread_All)] for r in cot.itertuples()]
ws=sheet('Futures positioning (CFTC)',['Report date (Tue)','CFTC contract code','Market (CFTC name)','Open interest, contracts','Managed money long','Managed money short','Managed money spreading'],cdat,
   [('Managed money NET long [formula]',lambda r:f'=E{r}-F{r}'),('Net long, % of open interest [formula]',lambda r:f'=(E{r}-F{r})/D{r}')],widths=[12,10,50,14,14,14,14,16,16],
   note='CFTC Commitments of Traders, Disaggregated Futures-Only (history files). 067651 = NYMEX WTI light sweet crude; 111659 = NYMEX RBOB gasoline. Disaggregated data begin Jun 13, 2006, so early 2008 is covered but not 2005-06.')
for row in ws.iter_rows(min_row=3,max_col=1):
    for c in row: c.number_format='yyyy-mm-dd'
n=20000; b=1; C0=11
ws.cell(b,C0,'Summary by year (formulas over the rows at left; new weeks appended by update.py are included)').font=Font(bold=True)
hdrs=['Year','Contract','Peak managed-money net long [formula]','Average net long [formula]','Average open interest [formula]']
for j,h in enumerate(hdrs): ws.cell(b+1,C0+j,h).font=Font(bold=True); ws.column_dimensions[L(C0+j)].width=18
k=b+2
for code in ['067651','111659']:
    for y in [2008,2022,2026]:
        ws.cell(k,C0,y); ws.cell(k,C0+1,code)
        cond=f'$B$3:$B${n},"{code}",$A$3:$A${n},">="&DATE({y},1,1),$A$3:$A${n},"<="&DATE({y},12,31)'
        ws.cell(k,C0+2,f'=_xlfn.MAXIFS($H$3:$H${n},{cond})'); ws.cell(k,C0+3,f'=AVERAGEIFS($H$3:$H${n},{cond})'); ws.cell(k,C0+4,f'=AVERAGEIFS($D$3:$D${n},{cond})'); k+=1
cot['net']=cot.M_Money_Positions_Long_All-cot.M_Money_Positions_Short_All; cot['pct']=cot.net/cot.Open_Interest_All
for code,g in cot.groupby('code'):
    g=g.set_index('date'); F[f'cot_{code}']={str(y):{'peak':int(g[g.index.year==y].net.max()),'peak_date':str(g[g.index.year==y].net.idxmax().date()),'avg':int(g[g.index.year==y].net.mean()),'peak_pct':round(float(g[g.index.year==y].pct.max())*100,1)} for y in [2008,2022,2026]}
    F[f'cot_{code}']['latest']={'date':str(g.index[-1].date()),'net':int(g.net.iloc[-1]),'pct':round(float(g.pct.iloc[-1])*100,1),'alltime_peak':int(g.net.max()),'alltime_peak_date':str(g.net.idxmax().date()),'rank_latest_pct':int((g.pct>g.pct.iloc[-1]).sum())+1,'n':len(g)}
# ---------------- investigations
INV=[('2006-05','FTC','Post-Katrina gasoline investigation','Finding','No evidence of illegal market manipulation; 15 firms met the statutory price-gouging definition, mostly explained by regional/market factors','https://www.ftc.gov/news-events/news/press-releases/2006/05/ftc-releases-report-its-investigation-gasoline-price-manipulation-post-katrina-gasoline-price'),
('2011-09','FTC Bureau of Economics','Gasoline price changes report','Finding (economic study)','Retail prices rise faster than they fall ("rockets and feathers"); causes not fully understood','https://www.ftc.gov/reports/federal-trade-commission-bureau-economics-gasoline-price-changes-petroleum-industry-update'),
('2012-04-19','CFTC / S.D.N.Y.','CFTC v. Optiver','Finding (consent order)','Manipulation/attempted manipulation of NYMEX crude, heating oil, NY Harbor gasoline futures in Mar 2007; $13M penalty + $1M disgorgement','https://www.cftc.gov/PressRoom/PressReleases/6239-12'),
('2024-02','Colorado Dept. of Public Health & Environment','Suncor Commerce City air violations','Finding (penalty) - pollution, not pricing','$10.5M penalty for three years of air-pollution violations (reported by Colorado Sun; CDPHE is the primary source). Unrelated to fuel prices','https://coloradosun.com/2026/09/10/suncor-colorado-new-rules-crackdown-air-pollution/'),
('2024-07-10','California AG','People v. Vitol & SK Energy','Settlement (not adjudicated)','$50M settlement over alleged 2015 California spot-gasoline manipulation; allegations settled, not proven','https://oag.ca.gov/news/press-releases/attorney-general-bonta-announces-50-million-settlement-vitol-and-sk-part-ongoing'),
('2025-08-29','California Energy Commission','SB X1-2 max refining margin','Decision (no penalty)','No margin cap or penalty adopted; action deferred at least 5 years; zero SB X1-2 penalty findings','https://efiling.energy.ca.gov/GetDocument.aspx?tn=265835'),
('2026-06-12','CEC Division of Petroleum Market Oversight (DPMO)','Station pricing inquiry','Statement / open inquiry (no finding)','~20 high-priced stations contacted after Mar 19 bulletin; no findings published; CA branded-unbranded spread $0.31 vs $0.06 elsewhere','https://www.energy.ca.gov/sites/default/files/2026-06/DPMO_California_Gasoline_Diesel_Market_Update_June_2026_ada.pdf'),
('2026-07-03','DOJ Antitrust Division + FTC','Letter to state AGs','Statement (no finding)','Says too much of crude price cut is being "withheld"; agencies "closely monitoring"; urges state probes','https://www.justice.gov/atr/media/1450951/dl?inline='),
('2026-07-03','DOJ Office of Public Affairs','Press release on the letter','Statement (no finding)','"closely monitoring petroleum markets"; no case, charge or subpoena announced','https://www.justice.gov/opa/pr/justice-department-and-federal-trade-commission-issue-call-action-state-attorneys-general'),
('2026 (through Sep 24)','Colorado Attorney General','Fuel price enforcement','None located','No Colorado AG investigation, complaint or finding on 2026 fuel prices was located. Colorado\'s documented 2026 Suncor actions (CDPHE) concern air pollution','https://coag.gov/'),
('2026 (through Sep 24)','CFTC','Energy futures manipulation, 2026','None located','No 2026 CFTC enforcement action on crude or RBOB gasoline manipulation located','https://www.cftc.gov/PressRoom/PressReleases'),
('2026 (through Sep 24)','FBI / DOJ','Subpoenas re "Democrats\' Big Oil payoff" (viral X post)','None located','No DOJ/FBI press release, court filing or other record of such subpoenas located. Claim rated Unsupported','https://www.justice.gov/news'),
('2026-09-24','Senate EPW minority (Democrats: Whitehouse, Schumer)','"The Billion Dollar Deal" report','Partisan staff report (allegations, not a finding)','Alleges oil & gas firms were rewarded after an Apr 2024 Mar-a-Lago $1B donation request (the request itself is reported, not confirmed by primary record)','https://www.epw.senate.gov/public/_cache/files/5/9/59c00f03-c375-4266-9bc9-51c1fb1ccc90/BD367324A2B0FF111B29E228E3C339AA1C8E84E703CED38A6644E80E040CB213.polluter-report-final.pdf')]
sheet('Investigations',['Date','Body','Matter','Type (finding vs statement)','What the record says','Source URL'],[list(x) for x in INV],widths=[14,30,30,26,80,70])
# ---------------- political money
os_rows=[]
import io
t=pd.read_html(io.StringIO(open(R+'os_E01_2024.html').read()))[0]
for r in t.itertuples(index=False):
    if int(r[0]) in (2020,2022,2024):
        os_rows.append([int(r[0]),r[1],r[5],r[6],r[7],r[8],'compiled from FEC filings by OpenSecrets (Oil & Gas industry, code E01); retrieved Sep 24 2026','https://www.opensecrets.org/industries/totals?cycle=2024&ind=E01'])
os_rows.append([2026,'not published','not published','not published','','','OpenSecrets industry totals table did not yet include the 2026 cycle on Sep 24 2026','https://www.opensecrets.org/industries/totals?cycle=2026&ind=E01'])
ws=sheet('Political money',['Election cycle','Total contributions (as shown)','To Democrats','To Republicans','% to Democrats','% to Republicans','Basis','Source URL'],os_rows,widths=[10,18,16,16,12,12,60,60],
   note='Part A: oil & gas industry totals, compiled from FEC filings by OpenSecrets (strings typed exactly as shown). Party split excludes outside/soft money not attributed to a party. Part B (sheet "FEC refiner PACs"): primary FEC bulk data.')
PAC={'C00109546':'Valero Energy Corp PAC','C00496307':'Marathon Petroleum Corp Employees PAC (MPAC)','C00513549':'Phillips 66 PAC','C00342766':'HF Sinclair PAC (DINO PAC; formerly HollyFrontier PAC)','C00415026':'American Fuel & Petrochemical Manufacturers PAC (AFPM PAC)','C00483677':'American Petroleum Institute PAC (API PAC)'}
fec=[]; cols='CMTE_ID AMNDT_IND RPT_TP TRANSACTION_PGI IMAGE_NUM TRANSACTION_TP ENTITY_TP NAME CITY STATE ZIP EMPLOYER OCCUPATION TRANSACTION_DT TRANSACTION_AMT OTHER_ID CAND_ID TRAN_ID FILE_NUM MEMO_CD MEMO_TEXT SUB_ID'.split()
for y in ['20','22','24','26']:
    cn=pd.read_csv(R+f'fec/cn{y}/cn.txt',sep='|',header=None,dtype=str,usecols=[0,1,2]); party=dict(zip(cn[0],cn[2]))
    p=pd.read_csv(R+f'fec/pas2{y}/itpas2.txt',sep='|',header=None,names=cols,dtype=str)
    p=p[p.CMTE_ID.isin(PAC)&(p.TRANSACTION_TP=='24K')]
    for r in p.itertuples():
        fec.append(['20'+y,r.CMTE_ID,PAC[r.CMTE_ID],r.NAME,r.CAND_ID,party.get(r.CAND_ID,''),r.TRANSACTION_DT,float(r.TRANSACTION_AMT),r.IMAGE_NUM,f'https://docquery.fec.gov/cgi-bin/fecimg/?{r.IMAGE_NUM}'])
fec.sort(key=lambda x:(x[0],x[2],x[6] or ''))
ws=sheet('FEC refiner PACs',['Cycle','FEC committee ID','PAC','Recipient (as filed)','Candidate ID','Candidate party (FEC candidate master)','Date (MMDDYYYY)','Amount, $ (as filed)','FEC image no.','Filing image URL'],fec,widths=[8,12,40,36,12,10,12,12,20,50],
   note='Source: FEC bulk files pas2 (contributions from committees to candidates, type 24K direct contributions) and cn (candidate master), cycles 2020-2026 (2026 partial). https://www.fec.gov/data/browse-data/?tab=bulk-data . No FEC-registered PAC found for PBF Energy or Suncor.')
n=ws.max_row; b=n+2; ws.cell(b,1,'Summary (formulas over rows above)').font=Font(bold=True)
for j,h in enumerate(['Cycle','To DEM, $ [formula]','To REP, $ [formula]','Other/none, $ [formula]','DEM share [formula]','REP share [formula]'],1): ws.cell(b+1,j,h).font=Font(bold=True)
for i,y in enumerate(['2020','2022','2024','2026']):
    r=b+2+i; ws.cell(r,1,y)
    ws.cell(r,2,f'=SUMIFS($H$3:$H${n},$A$3:$A${n},"{y}",$F$3:$F${n},"DEM")'); ws.cell(r,3,f'=SUMIFS($H$3:$H${n},$A$3:$A${n},"{y}",$F$3:$F${n},"REP")')
    ws.cell(r,4,f'=SUMIFS($H$3:$H${n},$A$3:$A${n},"{y}")-B{r}-C{r}'); ws.cell(r,5,f'=B{r}/(B{r}+C{r}+D{r})'); ws.cell(r,6,f'=C{r}/(B{r}+C{r}+D{r})')
fd=pd.DataFrame(fec,columns=['cycle','cid','pac','name','cand','party','date','amt','img','url'])
F['fec']={c:{p:round(float(g[g.party==p].amt.sum())) for p in ['DEM','REP']}|{'total':round(float(g.amt.sum())),'n':len(g)} for c,g in fd.groupby('cycle')}
F['fec_by_pac']={f'{c}|{p}':{q:round(float(h[h.party==q].amt.sum())) for q in ['DEM','REP']} for (c,p),h in fd.groupby(['cycle','pac'])}
F['opensecrets']=[r[:6] for r in os_rows]
# ---------------- facts for report & red flags
def at(s,d): return float(s.loc[pd.Timestamp(d)])
last=wdf.iloc[-1]; F['latest']={k:(float(last[k]) if not isinstance(last[k],(dt.date,)) else str(last[k])) for k in ['us','co','p4','ca','wti','brent','gc','ny','la','tax','gap','crack','crackb','ret_whl']}|{'week':str(last.week),'spot_week':str(last.spot_week)}
j08=wdf[wdf.week==dt.date(2008,7,7)].iloc[0]; F['jul7_2008']={k:float(j08[k]) for k in ['us','co','p4','ca','wti','brent','gc','ny','la','tax','gap','crack','crackb','ret_whl']}
wdf['year']=[d.year for d in wdf.week]
F['gap_annual']={int(y):round(float(g.gap.mean()),3) for y,g in wdf.groupby('year')}
F['crack_annual']={int(y):round(float(g.crack.mean()),3) for y,g in wdf.groupby('year')}
F['retwhl_annual']={int(y):round(float(g.ret_whl.mean()),3) for y,g in wdf.groupby('year')}
F['co_minus_us_annual']={int(y):round(float((g.co-g.us).mean()),3) for y,g in wdf.groupby('year')}
F['co_minus_p4_annual']={int(y):round(float((g.co-g.p4).mean()),3) for y,g in wdf.groupby('year')}
F['max_gap']={'val':float(wdf.gap.max()),'week':str(wdf.loc[wdf.gap.idxmax(),'week'])}
F['max_crack']={'val':float(wdf.crack.max()),'week':str(wdf.loc[wdf.crack.idxmax(),'week'])}
F['gap_latest_rank']=int((wdf.gap>last.gap).sum())+1; F['n_weeks']=len(wdf)
pc['ref_c']=pc.retail_usd_gal*pc.refining_pct; pc['dm_c']=pc.retail_usd_gal*pc.dist_marketing_pct; pc['tax_c']=pc.retail_usd_gal*pc.taxes_pct; pc['crude_c']=pc.retail_usd_gal*pc.crude_pct
F['ref_c_top']=[(str(r.month.date())[:7],round(r.ref_c,1)) for r in pc.nlargest(8,'ref_c').itertuples()]
F['ref_c_may26_rank']=int((pc.ref_c>pc[pc.month=='2026-05-01'].ref_c.iloc[0]).sum())+1
F['ref_c_2022_max']=[(str(r.month.date())[:7],round(r.ref_c,1)) for r in pc[pc.month.dt.year==2022].nlargest(1,'ref_c').itertuples()]
F['dm_c_top']=[(str(r.month.date())[:7],round(r.dm_c,1)) for r in pc.nlargest(5,'dm_c').itertuples()]
F['dm_c_2026']=[(str(r.month.date())[:7],round(r.dm_c,1),round(r.ref_c,1)) for r in pc[pc.month.dt.year==2026].itertuples()]
# pass-through windows (weekly)
def win(a,b):
    A=wdf[wdf.week==pd.Timestamp(a).date()].iloc[0]; B=wdf[wdf.week==pd.Timestamp(b).date()].iloc[0]
    return {'from':a,'to':b,'retail_chg':round(B.us-A.us,3),'gc_chg':round(B.gc-A.gc,3),'brent_chg_gal':round((B.brent-A.brent)/42,3),'wti_chg_gal':round((B.wti-A.wti)/42,3)}
F['win_down']=win('2026-06-01','2026-07-06'); F['win_up']=win('2026-07-06','2026-09-21'); F['win_war']=win('2026-02-23','2026-04-06')
_st=E['stocks']; _ld=_st.index[-1]; _same={}
for y in range(2005,_ld.year+1):
    t=pd.Timestamp(y,_ld.month,_ld.day); x=_st[(_st.index>=t-pd.Timedelta(days=3))&(_st.index<=t+pd.Timedelta(days=3))]
    if len(x): _same[y]=float(x.iloc[0])
_lower=[y for y,v in _same.items() if v<_same[_ld.year] and y<_ld.year]
F['stocks']={'latest':float(_st.iloc[-1]),'date':str(_ld.date()),'same_week_by_year':_same,'last_year_lower':max(_lower) if _lower else None}
# decomposition of today's refining margin vs WTI (all inputs EIA weekly)
bw25=float((W['brent']-W['wti'])['2025'].mean())/42; cb25=float((W['gc']-W['brent']/42)['2025'].mean()); cw25=float((W['gc']-W['wti']/42)['2025'].mean())
F['decomp']={'crude_wti':last.wti/42,'refining_vs_wti':last.crack,'taxes':last.tax,'dist_mkt':last.ret_whl-last.tax,'retail':last.us,
  'baseline_2025_gc_wti':cw25,'bw_today':(last.brent-last.wti)/42,'bw_2025':bw25,'bw_widening':(last.brent-last.wti)/42-bw25,'crackb_today':last.crackb,'crackb_2025':cb25,'crackb_excess':last.crackb-cb25,
  'j08_crude':j08.wti/42,'j08_refining':j08.crack,'j08_taxes':j08.tax,'j08_dm':j08.ret_whl-j08.tax,'j08_retail':j08.us}
_c=cot[cot.code=='111659'].set_index('date'); F['rbob_sep8_highest_since']=str(_c[(_c.net>=int(_c.loc['2026-09-08'].net))&(_c.index<'2026-09-08')].index.max().date()) if len(_c[(_c.net>=int(_c.loc['2026-09-08'].net))&(_c.index<'2026-09-08')]) else 'none'
_r=wdf.set_index('week'); F['retwhl_jun1']=float(_r.loc[dt.date(2026,6,1)].ret_whl); F['retwhl_jul6']=float(_r.loc[dt.date(2026,7,6)].ret_whl)
F['crack_max_prev']={'val':float(wdf[wdf.year<2026].crack.max()),'week':str(wdf.loc[wdf[wdf.year<2026].crack.idxmax(),'week'])}
gx=E['gasx']; F['gasx']={'war_avg':round(float(gx['2026-03-01':].mean())),'prewar_avg':round(float(gx['2025-12-01':'2026-02-27'].mean())),'y2025_avg':round(float(gx['2025'].mean())),'latest':float(gx.iloc[-1])}
F['refiners']=rdf.assign(b=rdf.val/1e9).pivot_table(index=['co','period'],columns='metric',values='b').round(3).reset_index().to_dict('records')
json.dump(F,open(G+'facts.json','w'),indent=1,default=str)
# ---------------- Sources tab
src=[('Weekly 2005-present',h,DN.format(s,'W')) for h,_,s in SER]+[('Weekly 2005-present','EIA pump-components tax share','https://www.eia.gov/petroleum/gasdiesel/gaspump_hist.php'),('Weekly 2005-present','EIA avg federal+state gasoline tax, Jul 2026 (52.16 c/gal)','https://www.eia.gov/petroleum/marketing/monthly/xls/fueltaxes.xlsx'),
 ('Monthly pump components','EIA Gasoline Pump Components History (all columns)','https://www.eia.gov/petroleum/gasdiesel/gaspump_hist.php'),('Monthly pump components','EIA methodology','https://www.eia.gov/petroleum/gasdiesel/pump_methodology.php'),
 ('Refiners (SEC)','Each row: SEC filing index URL in column I; company facts API','https://data.sec.gov/api/xbrl/companyfacts/CIK0001035002.json'),
 ('Ownership concentration','EIA Refinery Capacity Report 2026, Table 5 (capacity by corporation/plant)',T5),('Ownership concentration','EIA Refinery Capacity Report 2026, Table 3 (by state; Colorado = Suncor only)','https://www.eia.gov/petroleum/refinerycapacity/table3.pdf'),('Ownership concentration','EIA PADD definitions','https://www.eia.gov/petroleum/marketing/prime/pdf/padddef.pdf'),
 ('Ownership concentration','Colorado product pipelines (Magellan Chase from Kansas; three Texas Panhandle lines; Rocky Mountain System and Medicine Bow from Wyoming) - SECONDARY: API Colorado report (2018). Reported, not confirmed by primary record','https://www.api.org/~/media/files/news/2018/18-june/colorado_naturalgas_report-june-2018.pdf'),
 ('Exports vs inventories','EIA weekly gasoline stocks WGTSTUS1',DN.format('WGTSTUS1','W')),('Exports vs inventories','EIA weekly total motor gasoline exports',DN.format('W_EPM0F_EEX_NUS-Z00_MBBLD','W')),('Exports vs inventories','EIA weekly product exports WRPEXUS2',DN.format('WRPEXUS2','W')),('Exports vs inventories','EIA weekly crude exports WCREXUS2',DN.format('WCREXUS2','W')),('Exports vs inventories','EIA annual MGFEXUS2 / MTPEXUS2',DN.format('MGFEXUS2','A')),
 ('Futures positioning (CFTC)','CFTC COT disaggregated futures-only history files (2006-2016, 2022, 2026)','https://www.cftc.gov/MarketReports/CommitmentsofTraders/HistoricalCompressed/index.htm'),
 ('Investigations','Row-level URLs in column F',''),('Political money','OpenSecrets Oil & Gas industry totals - compiled from FEC filings by OpenSecrets','https://www.opensecrets.org/industries/totals?cycle=2024&ind=E01'),
 ('FEC refiner PACs','FEC bulk data: pas2 + cn files 2020-2026','https://www.fec.gov/data/browse-data/?tab=bulk-data'),
 ('Context','IEA Oil Market Report Sep 2026 (record Atlantic Basin margins; Gulf product exports down ~60%)','https://www.iea.org/reports/oil-market-report-september-2026'),('Context','EIA STEO Sep 2026','https://www.eia.gov/outlooks/steo/'),
 ('Context','EIA operable refinery capacity MOCLEUS2',DN.format('MOCLEUS2','M')),('Context','CEC FAQ on California prices and Iran','https://www.energy.ca.gov/programs-and-topics/topics/petroleum-fuels-transition/faq-california-gas-prices-and-impacts-iran')]
RF=json.load(open(G+'redflags.json'))
sheet('Unexplained - Red flags',['#','Data point (fixed values)','Value(s)','Documented explanation, if any','Status','Primary source URL(s)'],[[i+1]+r for i,r in enumerate(RF)],widths=[4,50,50,60,34,70])
sheet('Sources',['Tab','Series / number','URL'],[list(s) for s in src],widths=[26,90,100])
wb.save(G+'gas-gap-tracker.xlsx')
# ---------------- charts
plt.rcParams.update({'font.size':10,'text.parse_math':False})
w=wdf.copy(); w['d']=pd.to_datetime(w.week)
fig,ax=plt.subplots(figsize=(11,5.8))
ax.plot(w.d,w.us,color='#b91c1c',lw=1.6,label='U.S. retail regular (pump)')
ax.plot(w.d,w.gc,color='#1e3a5f',lw=1,label='Gulf Coast wholesale gasoline (spot)')
ax.plot(w.d,w.wti/42,color='#57534e',lw=1,label='WTI crude per gallon ($/bbl ÷ 42)')
ax.plot(w.d,w.brent/42,color='#d97706',lw=0.8,alpha=0.8,label='Brent crude per gallon')
ax.set_ylabel('Dollars per gallon'); ax.grid(alpha=0.3); ax.legend(loc='upper left',fontsize=8.5)
ax.set_title('Pump vs. wholesale gasoline vs. crude, weekly, Jan 2005 - Sep 2026 (nominal $/gal)')
ax.xaxis.set_major_locator(mdates.YearLocator(2)); ax.xaxis.set_major_formatter(mdates.DateFormatter('%Y'))
fig.text(0.01,0.01,'Source: EIA weekly series EMM_EPMR_PTE_NUS_DPG, EER_EPMRU_PF4_RGC_DPG, RWTC, RBRTE. Latest: retail %s, spot week ending %s.'%(last.week,last.spot_week),fontsize=7)
plt.tight_layout(rect=(0,0.03,1,1)); plt.savefig(G+'charts/gap-pump-wholesale-crude-2005-2026.png',dpi=140); plt.close()
p2=pc[pc.month>='2005-01-01']
fig,ax=plt.subplots(figsize=(11,5.8))
ax.stackplot(p2.month,p2.crude_c,p2.ref_c,p2.dm_c,p2.tax_c,colors=['#57534e','#b91c1c','#1e3a5f','#9ca3af'],labels=['Crude oil','Refining costs & profits','Distribution & marketing','Taxes'])
ax.set_ylabel('Cents per gallon (nominal)'); ax.legend(loc='upper left',fontsize=8.5); ax.grid(alpha=0.3)
ax.set_title('Where each gallon\'s price goes, monthly, Jan 2005 - May 2026 (EIA pump components)')
ax.xaxis.set_major_locator(mdates.YearLocator(2)); ax.xaxis.set_major_formatter(mdates.DateFormatter('%Y'))
fig.text(0.01,0.01,'Source: EIA Gasoline Pump Components History (EIA % share x EIA retail price). May 2026 is the latest month EIA has published.',fontsize=7)
plt.tight_layout(rect=(0,0.03,1,1)); plt.savefig(G+'charts/gap-cents-per-gallon-breakdown-2005-2026.png',dpi=140); plt.close()
ni=rdf[rdf.metric.str.startswith('Net income attributable')&rdf.co.isin(['Valero Energy','Marathon Petroleum','Phillips 66','PBF Energy','HollyFrontier (to 2021)','HF Sinclair (2022-)'])]
yrs=[str(y) for y in range(2012,2026)]; tot=[ni[ni.period==y].val.sum()/1e9 for y in yrs]; h1=ni[ni.period=='2026 H1 (Jan-Jun)'].val.sum()/1e9
F['ni5']=dict(zip(yrs,[round(x,3) for x in tot]))|{'2026H1':round(h1,3)}; json.dump(F,open(G+'facts.json','w'),indent=1,default=str)
cr=[F['crack_annual'][int(y)] for y in yrs]
fig,ax=plt.subplots(figsize=(11,5.6)); ax2=ax.twinx()
ax.bar(yrs+['2026\nH1 only'],tot+[h1],color=['#1e3a5f']*len(yrs)+['#b91c1c'])
ax2.plot(yrs+['2026\nH1 only'],cr+[float(wdf[wdf.year==2026].crack.mean())],color='#d97706',marker='o',lw=2,label='Gulf Coast gasoline minus WTI, $/gal (annual avg; 2026 Jan-Sep)')
ax.set_ylabel('Combined net income, $ billions (SEC)'); ax2.set_ylabel('Crack spread, $/gal'); ax2.set_ylim(0,max(cr)*1.6)
ax.set_title('Refiner profits vs. refining margin: Valero, Marathon Petroleum, Phillips 66, PBF, HF Sinclair/HollyFrontier')
ax2.legend(loc='upper left',fontsize=8.5); ax.grid(alpha=0.3,axis='y')
fig.text(0.01,0.01,'Sources: SEC XBRL NetIncomeLoss (10-K; 2026 = Jan-Jun from Q2 10-Qs, half a year); EIA weekly spot prices. MPC and PSX include non-refining segments.',fontsize=7)
plt.tight_layout(rect=(0,0.03,1,1)); plt.savefig(G+'charts/gap-refiner-profits-vs-margin.png',dpi=140); plt.close()
print(json.dumps({k:F[k] for k in F if k not in('refiners','fec_by_pac','gap_annual','crack_annual','retwhl_annual','co_minus_us_annual','co_minus_p4_annual')},indent=0,default=str)[:6000])
